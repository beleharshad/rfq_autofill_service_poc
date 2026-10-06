"""Real CAD integration with controlled vision responses; not live AI accuracy."""
import math
import sys
import cadquery as cq
import fitz
import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.api import part_spec as routes
from app.models.automatic import DrawingProposal, DrawingAudit
from app.models.document_registry import AssociationRequest
from app.services.drawing_interpreter import interpret_drawing, normalize_proposal, request_json, ProviderUnavailable
from app.services.automatic_workflow import execute
from app.services.automatic_runner import run_automatic
from app.storage.automatic_runs import AutomaticRuns
from app.storage.accepted_parts import AcceptedParts
from app.storage.file_storage import FileStorage
from app.security import require_api_key
from app.workers.generic_queue import drain
JOB='aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee'


def pdf_bytes():
    doc=fitz.open();page=doc.new_page();page.insert_text((30,30),'Synthetic regression drawing')
    data=doc.tobytes();doc.close();return data


def drawing(part='TEST',units='in'):
    s=1 if units=='in' else 25.4
    p=DrawingProposal.model_validate(dict(part_number=part,revision='A',role='finished_drawing',units=units,
        base=dict(kind='cylinder',diameter=2*s,length=s),
        features=[dict(kind='hole',feature_id='bore',origin=[0,0,0],axis=[0,0,1],diameter=.5*s,depth=s,termination='through')],
        complete=True,unresolved=[],evidence=[dict(target='base',page=1,callout='OD 2 in, length 1 in'),dict(target='bore',page=1,callout='Diameter .5 in through')]))
    a=DrawingAudit.model_validate(dict(agrees=True,part_number=part,revision='A',finished_part=True,
        all_features_accounted_for=True,checked_targets=['base','bore'],issues=[],dimensions=[
        dict(metric='axial_length_mm',lower=25.4,upper=25.4,page=1,callout='1 in'),
        dict(metric='maximum_outer_cylindrical_diameter_mm',lower=50.8,upper=50.8,page=1,callout='2 in'),
        dict(metric='minimum_coaxial_through_bore_mm',lower=12.7,upper=12.7,page=1,callout='.5 in')]))
    return p,a,1


def prepare_run(tmp_path,files=None):
    storage=FileStorage(tmp_path);storage.read_base_paths=[tmp_path]
    for name,data in files or [('source.pdf',pdf_bytes())]:storage.save_bytes_file(JOB,name,data)
    path=storage.get_job_path(JOB);runs=AutomaticRuns(path);run=runs.enqueue();assert runs.claim(run['run_id'])
    return storage,path,runs,run


def test_positioned_holes_threads_and_blind_bottoms_reach_export(tmp_path):
    """Controlled retainer-like recipe; tests integration, not drawing interpretation.

    Representative grooves are explicit test inputs, not certified UNC dimensions.
    This simplified fixture intentionally does not claim to model drawing fillets.
    """
    _, path, runs, run = prepare_run(tmp_path)
    features = []
    centers = [(17.526, 0), (0, 17.526), (-17.526, 0), (0, -17.526)]
    for i, (x, y) in enumerate(centers):
        features.append(dict(kind='hole', feature_id=f'hole{i}', origin=[x,y,0],
                             axis=[0,0,1], diameter=11.0236, depth=25.4, termination='through'))
        features.append(dict(kind='thread', feature_id=f'thread{i}', origin=[x,y,0],
                             axis=[0,0,1], side='internal', major_diameter=12.7,
                             minor_diameter=11.0236, pitch=25.4/13, length=25.4,
                             groove_width=.8, callout='Representative 0.500-13 UNC-2B'))
    for i, x in enumerate([-8,8]):
        features.append(dict(kind='hole', feature_id=f'blind{i}', origin=[x,0,25.4],
                             axis=[0,0,-1], diameter=5.2578, depth=16.764, termination='blind'))
    p = DrawingProposal.model_validate(dict(part_number='POSITIONED-REGRESSION',revision='A',
        role='finished_drawing', units='mm', complete=True, unresolved=[],
        base=dict(kind='revolve', profile=[[0,0],[28.3845,0],[28.3845,11.176],
                   [27.2415,11.176],[27.2415,25.4],[0,25.4]]), features=features,
        evidence=[dict(target=t,page=1,callout='Explicit synthetic test geometry')
                  for t in ['base']+[f['feature_id'] for f in features]]))
    audit = DrawingAudit.model_validate(dict(agrees=True,part_number=p.part_number,revision='A',
        finished_part=True,all_features_accounted_for=True,issues=[],
        checked_targets=['base']+[f['feature_id'] for f in features],dimensions=[
            dict(metric='axial_length_mm',lower=25.399,upper=25.401,page=1,callout='test length'),
            dict(metric='maximum_outer_cylindrical_diameter_mm',lower=56.768,upper=56.770,page=1,callout='test OD')]))
    execute(path,run['run_id'],interpreter=lambda source:(p,audit,1))
    result=runs.get()
    assert result['status']=='complete',result
    item=result['result']['items'][0]
    output=path/'outputs'/'accepted'/item['build_id']
    solid=cq.importers.importStep(str(output/'model.step')).solids().val()
    assert solid.isValid()
    assert solid.isInside((0,0,12.7)) # no central bore
    for x,y in centers:
        for z in (.1,12.7,25.3):assert not solid.isInside((x,y,z))
    for x in (-8,8):
        assert not solid.isInside((x,0,20))
        assert solid.isInside((x,0,5)) # actual blind-hole floor
    assert (output/'model.glb').read_bytes()[:4]==b'glTF'
    assert sum('representative helical' in w for w in item['warnings'])==4


@pytest.mark.parametrize('units',['in','mm'])
def test_pdf_to_actual_hollow_solid_and_exports(tmp_path,units):
    files,path,runs,run=prepare_run(tmp_path)
    execute(path,run['run_id'],interpreter=lambda source:drawing(units=units))
    result=runs.get();assert result['status']=='complete',result
    item=result['result']['items'][0];build=AcceptedParts(path).build(item['build_id'])
    assert build['status']=='complete' and build['manifest']['acceptance_origin']=='automatic_drawing'
    assert build['manifest']['status']=='constructed_automatically'
    assert build['manifest']['measurements']['volume_mm3']==pytest.approx(math.pi*(25.4**2-6.35**2)*25.4)
    output=path/'outputs'/'accepted'/item['build_id']
    solid=cq.importers.importStep(str(output/'model.step')).solids().val()
    assert not solid.isInside((0,0,12.7)) and solid.isInside((20,0,12.7))
    assert (output/'model.glb').read_bytes()[:4]==b'glTF'
    wb=load_workbook(output/'rfq_review.xlsx',data_only=True)
    assert any(r[1]=='automatic_drawing' for r in wb.active.iter_rows(values_only=True))
    assert runs.enqueue()['run_id']==run['run_id']


@pytest.mark.parametrize('failure',['omission','dimension','coverage','identity','units','duplicate','page'])
def test_uncertain_drawings_never_publish(tmp_path,failure):
    _,path,runs,run=prepare_run(tmp_path);p,a,n=drawing()
    if failure=='omission':a.all_features_accounted_for=False
    if failure=='dimension':a.dimensions[0].lower=30;a.dimensions[0].upper=31
    if failure=='coverage':p.evidence=[]
    if failure=='identity':a.revision='B'
    if failure=='units':p.units='mm'
    if failure=='duplicate':p.features.append(p.features[0])
    if failure=='page':p.evidence[0].page=2
    execute(path,run['run_id'],interpreter=lambda source:(p,a,n))
    result=runs.get();assert result['status']=='review_required'
    item=result['result']['items'][0];assert item['issues']
    if item.get('build_id'):
        assert AcceptedParts(path).build(item['build_id'])['status']=='failed'
        assert not (path/'outputs'/'accepted'/item['build_id']/'model.glb').exists()
    else:assert not AcceptedParts(path).list()


def test_batch_associations_keep_sibling_models_current(tmp_path):
    _,path,runs,run=prepare_run(tmp_path,[('one.pdf',pdf_bytes()),('two.pdf',pdf_bytes())]);names=iter(['ONE','TWO'])
    execute(path,run['run_id'],interpreter=lambda source:drawing(part=next(names)))
    assert runs.get()['status']=='complete'
    assert len(AcceptedParts(path).list())==2
    assert not any(s['stale'] for s in AcceptedParts(path).list())


def test_competing_sources_require_authority_choice(tmp_path):
    _,path,runs,run=prepare_run(tmp_path,[('one.pdf',pdf_bytes()),('two.pdf',pdf_bytes())])
    execute(path,run['run_id'],interpreter=lambda source:drawing())
    assert runs.get()['status']=='review_required' and not AcceptedParts(path).list()
    assert all(not d['association'] for d in runs.registry.read()['documents'])


def test_process_association_is_preserved(tmp_path):
    _,path,runs,run=prepare_run(tmp_path);doc=runs.registry.read()['documents'][0]
    changed=runs.registry.associate(doc['document_id'],AssociationRequest(expected_version=run['registry_version'],part_number='TEST',revision='A',role='process_drawing',manufacturing_state='intermediate'))
    runs.update(run['run_id'],registry_version=changed['version'])
    execute(path,run['run_id'],interpreter=lambda source:drawing())
    assert runs.get()['status']=='review_required' and not AcceptedParts(path).list()
    assert runs.registry.read()['documents'][0]['association']['role']=='process_drawing'


def step_bytes(tmp_path,multi=False):
    shape=cq.Workplane('XY').box(20,30,40).val()
    if multi:shape=cq.Compound.makeCompound([shape,shape.translate((100,0,0))])
    path=tmp_path/'fixture.step';cq.exporters.export(shape,str(path));return path.read_bytes()


def test_step_subprocess_and_persistent_queue(tmp_path,monkeypatch):
    files=FileStorage(tmp_path/'jobs');files.read_base_paths=[files.base_path]
    files.save_bytes_file(JOB,'solid.step',step_bytes(tmp_path))
    path=files.get_job_path(JOB);runs=AutomaticRuns(path);runs.enqueue()
    monkeypatch.setenv('GENERIC_GEOMETRY_PYTHON',sys.executable);drain(files.base_path)
    result=runs.get();assert result['status']=='complete',result
    item=result['result']['items'][0];assert item['measurements']['volume_mm3']==pytest.approx(24000)
    assert AcceptedParts(path).get(item['spec_id'])['payload']['acceptance_origin']=='source_cad'
    assert runs.registry.read()['documents'][0]['association']['manufacturing_state']=='unresolved'
    retry=runs.enqueue(retry=True);run_automatic(path,retry['run_id'])
    assert runs.get()['status']=='review_required' and len(AcceptedParts(path).list())==1


def test_multibody_step_not_silently_reduced(tmp_path):
    _,path,runs,run=prepare_run(tmp_path,[('assembly.step',step_bytes(tmp_path,True))]);execute(path,run['run_id'])
    assert runs.get()['status']=='review_required'
    assert '2 solids' in runs.get()['result']['items'][0]['issues'][0]
    assert not AcceptedParts(path).list()


def test_missing_credentials_durable_actionable_block(tmp_path,monkeypatch):
    monkeypatch.delenv('GOOGLE_API_KEY',raising=False)
    _,path,runs,run=prepare_run(tmp_path);execute(path,run['run_id'])
    item=runs.get()['result']['items'][0]
    assert item['status']=='blocked' and 'GOOGLE_API_KEY' in item['issues'][0]
    assert not AcceptedParts(path).list()


def test_two_pass_vision_renders_source(tmp_path):
    source=tmp_path/'scan.pdf';source.write_bytes(pdf_bytes());calls=[]
    def provider(pages,prompt,schema):
        assert pages[0][:8]==b'\x89PNG\r\n\x1a\n';calls.append(schema.__name__)
        return drawing()[0 if schema is DrawingProposal else 1]
    proposal,audit,pages=interpret_drawing(source,provider=provider)
    normalized,issues=normalize_proposal(proposal,audit,pages)
    assert not issues and normalized['base']['diameter']==50.8
    assert calls==['DrawingProposal','DrawingAudit']


def test_provider_http_contract_and_secret_redaction(monkeypatch):
    import requests
    monkeypatch.setenv('GOOGLE_API_KEY','secret-never-log');calls=[]
    class Response:
        status_code=200
        def json(self):return {'candidates':[{'finishReason':'STOP','content':{'parts':[{'text':drawing()[0].model_dump_json()}]}}]}
    def post(url,**kwargs):calls.append((url,kwargs));return Response()
    monkeypatch.setattr(requests,'post',post);request_json([b'png'],'prompt',DrawingProposal)
    assert 'secret' not in calls[0][0] and calls[0][1]['headers']['x-goog-api-key']=='secret-never-log'
    assert calls[0][1]['json']['contents'][0]['parts'][1]['inlineData']['mimeType']=='image/png'
    Response.status_code=429
    with pytest.raises(ProviderUnavailable,match='HTTP 429') as exc:request_json([b'png'],'prompt',DrawingProposal)
    assert 'secret-never-log' not in str(exc.value)


def test_source_change_prevents_acceptance(tmp_path):
    files,path,runs,run=prepare_run(tmp_path)
    def interpret(source):files.save_bytes_file(JOB,'new.pdf',pdf_bytes());return drawing()
    execute(path,run['run_id'],interpreter=interpret)
    assert runs.get()['status']=='failed' and not AcceptedParts(path).list()


def test_cancelled_run_cannot_publish(tmp_path):
    _,path,runs,run=prepare_run(tmp_path)
    def interpret(source):runs.cancel(run['run_id']);return drawing()
    execute(path,run['run_id'],interpreter=interpret)
    assert runs.get()['status']=='cancelled' and not AcceptedParts(path).list()


def test_automatic_api_auth_queue_and_artifacts(tmp_path,monkeypatch):
    files=FileStorage(tmp_path/'jobs');files.read_base_paths=[files.base_path]
    files.save_bytes_file(JOB,'solid.step',step_bytes(tmp_path))
    monkeypatch.setenv('INTERNAL_API_KEY','test-key');monkeypatch.setenv('GENERIC_GEOMETRY_PYTHON',sys.executable)
    monkeypatch.setenv('GENERIC_GEOMETRY_QUEUE_ONLY','true');monkeypatch.setattr(routes,'FileStorage',lambda:files)
    class Jobs:
        def get_job(self,job):return object() if job==JOB else None
    monkeypatch.setattr(routes,'JobStorage',Jobs)
    app=FastAPI();app.include_router(routes.router,prefix='/jobs',dependencies=[Depends(require_api_key)])
    client=TestClient(app);base='/jobs/'+JOB
    assert client.post(base+'/automatic').status_code==401
    client.headers['X-API-Key']='test-key'
    first=client.post(base+'/automatic');assert first.status_code==202
    assert client.post(base+'/automatic').json()['run_id']==first.json()['run_id']
    drain(files.base_path)
    done=client.get(base+'/automatic').json();assert done['status']=='complete'
    build=done['result']['items'][0]['build_id']
    for name in ['model.step','model.glb','rfq_review.xlsx','manifest.json']:
        assert client.get(base+f'/geometry-builds/{build}/artifacts/{name}').status_code==200
    files.save_bytes_file(JOB,'new.pdf',pdf_bytes())
    assert client.get(base+'/automatic').json()['stale']
    assert client.get(base+f'/geometry-builds/{build}/artifacts/model.glb').status_code==409


def test_real_multipart_upload_starts_automatic_workflow(tmp_path,monkeypatch):
    from app.api import jobs
    from app.services.job_service import JobService
    from app.storage.job_storage import JobStorage
    files=FileStorage(tmp_path/'jobs');files.read_base_paths=[files.base_path]
    job_store=JobStorage(tmp_path/'jobs.sqlite3');job_store.read_db_paths=[job_store.db_path]
    service=JobService();service.file_storage=files;service.job_storage=job_store
    # Legacy STEP preview is optional and independent of the new CAD worker.
    service.step_analysis_service.file_storage=files
    monkeypatch.setattr(jobs,'job_service',service)
    monkeypatch.setattr(routes,'FileStorage',lambda:files);monkeypatch.setattr(routes,'JobStorage',lambda:job_store)
    monkeypatch.setenv('INTERNAL_API_KEY','upload-key');monkeypatch.setenv('GENERIC_GEOMETRY_QUEUE_ONLY','true')
    monkeypatch.setenv('GENERIC_GEOMETRY_PYTHON',sys.executable)
    app=FastAPI()
    for router in (jobs.router,routes.router):app.include_router(router,prefix='/jobs',dependencies=[Depends(require_api_key)])
    client=TestClient(app,headers={'X-API-Key':'upload-key'})
    uploaded=client.post('/jobs',data={'mode':'auto_convert','automatic_geometry':'true'},files=[('files',('part.step',step_bytes(tmp_path),'application/step'))])
    assert uploaded.status_code==201,uploaded.text
    job=uploaded.json()['job_id'];base='/jobs/'+job
    assert client.get(base+'/automatic').json()['status']=='queued'
    drain(files.base_path)
    result=client.get(base+'/automatic').json();assert result['status']=='complete',result
    build=result['result']['items'][0]['build_id']
    assert client.get(base+f'/geometry-builds/{build}/artifacts/model.glb').content[:4]==b'glTF'
